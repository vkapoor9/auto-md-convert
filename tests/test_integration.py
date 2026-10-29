from __future__ import annotations

import contextlib
import importlib.util
import io
import json
import os
import plistlib
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import bootstrap
from dropzone import install_quick_action, watch
from shared.convert import core

REPO = Path(__file__).resolve().parents[1]
MARKER = "Audit fixture content"


def make_pdf(path: Path):
    stream = f"BT /F1 12 Tf 72 720 Td ({MARKER}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data)
    data += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    for offset in offsets[1:]:
        data += f"{offset:010} 00000 n \n".encode()
    data += f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)


def make_xls(path: Path):
    def record(code, content):
        return struct.pack("<HH", code, len(content)) + content
    text = MARKER.encode("ascii")
    path.write_bytes(
        record(0x0009, struct.pack("<HH", 2, 0x0010)) +
        record(0x0042, struct.pack("<H", 1252)) +
        record(0x0004, struct.pack("<HH", 0, 0) + b"\0\0\0" + bytes([len(text)]) + text) +
        record(0x000A, b"")
    )


class FormatIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def require(self, *names):
        for name in names:
            if importlib.util.find_spec(name) is None:
                self.skipTest(f"optional dependency {name} not installed")

    def assert_conversion(self, path):
        meta, body = core.convert_file(path)
        self.assertIn(MARKER, body)
        self.assertEqual(meta["type"], path.suffix[1:])
        self.assertEqual(meta["source_file"], path.name)
        output = core.write_markdown(path.parent, meta, body)
        self.assertTrue(output.read_text().startswith("---\n"))
        self.assertIn(MARKER, output.read_text())
        self.assertTrue(path.exists())

    def test_builtin_formats_without_optional_converters(self):
        for suffix, body in {
            ".txt": MARKER, ".csv": f"title,value\n{MARKER},1",
            ".json": json.dumps({"title": MARKER}),
            ".xml": f"<title>{MARKER}</title>",
            ".html": f"<html><body><p>{MARKER}</p></body></html>",
            ".htm": f"<html><body><p>{MARKER}</p></body></html>",
        }.items():
            with self.subTest(suffix=suffix):
                path = self.root / ("fixture" + suffix)
                path.write_text(body)
                with mock.patch.dict(sys.modules, {"markitdown": None, "html2text": None}):
                    self.assert_conversion(path)

    @unittest.skipUnless(shutil.which("textutil"), "macOS textutil required")
    def test_all_textutil_office_formats_without_markitdown(self):
        source = self.root / "source.txt"
        source.write_text(MARKER)
        for suffix in ("doc", "docx", "rtf", "odt"):
            with self.subTest(suffix=suffix):
                path = self.root / f"fixture.{suffix}"
                result = subprocess.run(
                    ["textutil", "-convert", suffix, "-output", str(path), str(source)],
                    capture_output=True, text=True, timeout=30,
                )
                self.assertEqual(result.returncode, 0, result.stderr)
                with mock.patch.dict(sys.modules, {"markitdown": None}):
                    self.assert_conversion(path)

    def test_real_pdf(self):
        self.require("markitdown", "pdfminer")
        path = self.root / "fixture.pdf"
        make_pdf(path)
        self.assert_conversion(path)

    def test_real_pptx(self):
        self.require("markitdown", "pptx")
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = MARKER
        path = self.root / "fixture.pptx"
        deck.save(path)
        self.assert_conversion(path)

    def test_real_xlsx_tables_and_blanks(self):
        self.require("openpyxl")
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Title", "Missing", "Escaped"])
        sheet.append([MARKER, None, "a|b\nline"])
        workbook.create_sheet("Empty")
        path = self.root / "fixture.xlsx"
        workbook.save(path)
        workbook.close()
        self.assert_conversion(path)
        _, body = core.convert_file(path)
        self.assertIn("a\\|b line", body)
        self.assertIn(f"| {MARKER} |  |", body)
        self.assertNotIn("NaN", body)
        self.assertNotIn("## Empty", body)

    def test_xlsx_preserves_formulas_without_cached_results(self):
        self.require("openpyxl")
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Total"])
        sheet.append(["=SUM(1,2)"])
        path = self.root / "formulas.xlsx"
        workbook.save(path)
        workbook.close()
        _, body = core.convert_file(path)
        self.assertIn("=SUM(1,2)", body)

    def test_real_legacy_xls(self):
        self.require("markitdown", "xlrd")
        path = self.root / "fixture.xls"
        make_xls(path)
        self.assert_conversion(path)

    def test_real_epub(self):
        self.require("ebooklib")
        from ebooklib import epub
        book = epub.EpubBook()
        book.set_identifier("audit-fixture")
        book.set_title("Audit fixture")
        book.set_language("en")
        chapter = epub.EpubHtml(title="Chapter", file_name="chapter.xhtml", lang="en")
        chapter.content = f"<html><body><p>{MARKER}</p></body></html>"
        book.add_item(chapter)
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.spine = ["nav", chapter]
        path = self.root / "fixture.epub"
        epub.write_epub(str(path), book)
        self.assert_conversion(path)

    def test_watcher_real_drop_then_same_name_replacement(self):
        src = self.root / "report.txt"
        src.write_text("First document")
        cfg = {"watchFolders": [str(self.root)], "pollSeconds": 1}
        seen = set()
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertTrue(watch._scan_once(cfg, seen))
            src.write_text("Second document")
            self.assertTrue(watch._scan_once(cfg, seen))
        self.assertEqual(len(list(self.root.glob("*.md"))), 2)
        self.assertEqual(sorted(p.read_text() for p in (self.root / ".originals").iterdir()),
                         ["First document", "Second document"])

    def test_live_watcher_process_converts_repeated_drops(self):
        app = self.root / "app"
        for directory in ("shared", "dropzone"):
            shutil.copytree(REPO / directory, app / directory,
                            ignore=shutil.ignore_patterns("__pycache__"))
        drops = self.root / "drops"
        drops.mkdir()
        (app / "config.json").write_text(json.dumps({
            "watchFolders": [str(drops)], "pollSeconds": 1,
        }))
        process = subprocess.Popen(
            [sys.executable, str(app / "dropzone/watch.py")],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
        )
        try:
            for number in (1, 2):
                source = drops / "report.txt"
                source.write_text(f"document {number}")
                deadline = time.monotonic() + 10
                archived = drops / ".originals" / ("report.txt" if number == 1 else "report-2.txt")
                while (not archived.exists() or source.exists()) and time.monotonic() < deadline:
                    self.assertIsNone(process.poll(), "watcher exited unexpectedly")
                    time.sleep(0.05)
                self.assertTrue(archived.exists(), "watcher did not respond to a new drop")
                self.assertFalse(source.exists(), "watcher did not finish archiving the drop")
                self.assertEqual(archived.read_text(), f"document {number}")
            self.assertIsNone(process.poll())
            self.assertEqual(len(list(drops.glob("*.md"))), 2)
        finally:
            process.terminate()
            try:
                stdout, stderr = process.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                stdout, stderr = process.communicate()
        self.assertEqual(stderr, "")
        self.assertIn("watching 1 folder", stdout)


class InstallIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_bootstrap_preserves_config_and_builds_valid_launch_agent(self):
        config = self.root / "config.json"
        config.write_text('{"watchFolders": [], "custom": "keep"}')
        env = {
            "AUTOMDCONVERT_SKIP_LOAD": "1", "AUTOMDCONVERT_SKIP_PIP": "1",
            "AUTOMDCONVERT_SKIP_OPEN": "1", "AUTOMDCONVERT_SKIP_QUICKACTION": "1",
            "AUTOMDCONVERT_LAUNCHAGENTS_DIR": str(self.root / "agents"),
            "AUTOMDCONVERT_ANSWERS": "", "SPLASH_OFF": "1",
        }
        with mock.patch.dict(os.environ, env), mock.patch.object(bootstrap, "REPO", self.root), \
             mock.patch.object(bootstrap, "_answers_iter", None), \
             contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(bootstrap.main(), 0)
        self.assertEqual(json.loads(config.read_text())["custom"], "keep")
        plist = self.root / "agents" / "com.automdconvert.watcher.plist"
        data = plistlib.loads(plist.read_bytes())
        self.assertEqual(data["ProgramArguments"], [sys.executable, str(self.root / "dropzone/watch.py")])
        self.assertTrue(data["KeepAlive"])

    def test_bootstrap_refuses_corrupt_config(self):
        config = self.root / "config.json"
        config.write_text("{broken")
        with mock.patch.object(bootstrap, "REPO", self.root), \
             mock.patch.object(bootstrap, "ask") as ask, \
             contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(bootstrap.main(), 1)
        ask.assert_not_called()
        self.assertEqual(config.read_text(), "{broken")

    def test_failed_launch_agent_is_not_reported_as_complete(self):
        env = {
            "AUTOMDCONVERT_SKIP_PIP": "1", "AUTOMDCONVERT_SKIP_OPEN": "1",
            "AUTOMDCONVERT_SKIP_QUICKACTION": "1", "AUTOMDCONVERT_ANSWERS": "",
            "AUTOMDCONVERT_LAUNCHAGENTS_DIR": str(self.root / "agents"),
        }
        output = io.StringIO()
        with mock.patch.dict(os.environ, env), mock.patch.object(bootstrap, "REPO", self.root), \
             mock.patch.object(bootstrap, "_answers_iter", None), \
             mock.patch.object(bootstrap, "load_launch_agent", return_value=False), \
             contextlib.redirect_stdout(output):
            self.assertEqual(bootstrap.main(), 1)
        self.assertIn("Setup incomplete", output.getvalue())
        self.assertNotIn("Setup complete.", output.getvalue())

    @unittest.skipUnless(shutil.which("plutil"), "macOS plutil required")
    def test_quick_action_plists_pass_real_validation(self):
        with mock.patch.object(install_quick_action, "services_dir", return_value=self.root), \
             mock.patch.object(install_quick_action, "_refresh_services"):
            self.assertTrue(install_quick_action.install(python=sys.executable))
            self.assertTrue(install_quick_action.uninstall())
        self.assertFalse((self.root / "Convert to Markdown.workflow").exists())

    @unittest.skipUnless(shutil.which("ditto"), "macOS ditto required")
    def test_package_includes_only_runtime_files_and_works_without_terminal(self):
        repo = self.root / "repo"
        repo.mkdir()
        for name in ("README.md", "LICENSE", "bootstrap.py", "config.example.json", "handoff-prompt.md",
                     "setup.command", "package.command", "shortcuts.zsh", "splash.py", ".gitignore"):
            shutil.copy2(REPO / name, repo / name)
        for name in ("shared", "dropzone"):
            shutil.copytree(REPO / name, repo / name, ignore=shutil.ignore_patterns("__pycache__"))
        for name in ("config.json", "private-notes.txt", "private.pdf", ".env"):
            (repo / name).write_text("DO NOT DISTRIBUTE")
        (repo / ".venv").mkdir()
        (repo / ".venv" / "private").write_text("DO NOT DISTRIBUTE")
        home = self.root / "home"
        home.mkdir()
        result = subprocess.run(["bash", str(repo / "package.command")], input="",
                                capture_output=True, text=True,
                                env={**os.environ, "HOME": str(home)}, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        archives = list((home / "Desktop").glob("*.zip"))
        self.assertEqual(len(archives), 1)
        with zipfile.ZipFile(archives[0]) as archive:
            names = archive.namelist()
            self.assertIn("auto-md-convert/setup.command", names)
            self.assertIn("auto-md-convert/LICENSE", names)
            self.assertIn("auto-md-convert/shared/convert/core.py", names)
            self.assertFalse(any("private" in n or "/.venv/" in n or n.endswith("config.json")
                                 or n.endswith(".env") for n in names))
            self.assertEqual(archive.getinfo("auto-md-convert/setup.command").external_attr >> 16 & 0o111,
                             0o111)
            archive.extractall(self.root / "unpacked")
        packaged = self.root / "unpacked/auto-md-convert"
        drops = self.root / "drops"
        drops.mkdir()
        source = drops / "document.txt"
        source.write_text(MARKER)
        (packaged / "config.json").write_text(json.dumps({"watchFolders": [str(drops)]}))
        result = subprocess.run(
            [sys.executable, str(packaged / "dropzone/watch.py"), "--once"],
            capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertFalse(source.exists())
        self.assertEqual((drops / ".originals/document.txt").read_text(), MARKER)
        self.assertIn(MARKER, next(drops.glob("*.md")).read_text())

    def test_setup_launcher_propagates_bootstrap_failure(self):
        shutil.copy2(REPO / "setup.command", self.root / "setup.command")
        (self.root / "bootstrap.py").write_text("raise SystemExit(7)\n")
        (self.root / "splash.py").write_text("")
        result = subprocess.run(["bash", str(self.root / "setup.command")], input="",
                                capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 7, result.stdout + result.stderr)


if __name__ == "__main__":
    unittest.main()
