"""Convert dropped files into clean Markdown a coding agent can read.

Dispatch by extension:
  - .eml            parsed here into email Markdown (headers, full thread, and
                    recursively converted attachments).
  - text / markup   built-in readers, with optional html2text for HTML.
  - spreadsheets    openpyxl first for XLSX, markitdown for legacy XLS.
  - other documents Microsoft markitdown or macOS `textutil` for Word / RTF / ODT.

Baseline is standard library plus macOS `textutil`; markitdown and html2text are
optional upgrades. Used by dropzone/watch.py.
"""

from __future__ import annotations

import json
import errno
import logging
import math
import mimetypes
import os
import re
import html as _html
import shutil
import subprocess
import tempfile
from collections.abc import Iterator
from contextlib import closing
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = REPO_ROOT / "config.json"

_DEFAULT_CONFIG = {
    "watchFolders": [],
    "pollSeconds": 2,
}

# Files the watcher will try to convert. Markdown output (.md) is intentionally
# excluded so the tool never reprocesses its own output.
SUPPORTED_EXTS = {
    ".eml",
    ".doc", ".docx", ".rtf", ".odt",
    ".pptx",
    ".xls", ".xlsx", ".csv",
    ".pdf",
    ".html", ".htm",
    ".epub",
    ".json", ".xml",
    ".txt",
}

# Formats macOS `textutil` can read, used when markitdown is unavailable.
_TEXTUTIL_EXTS = {".doc", ".docx", ".rtf", ".odt", ".html", ".htm", ".wordml", ".webarchive"}

# Formats markitdown handles poorly (returns raw markup); prefer textutil for these.
_TEXTUTIL_FIRST = {".rtf", ".doc", ".odt", ".wordml", ".webarchive"}

# Spreadsheets we convert with openpyxl first: cleaner tables, blanks instead of
# markitdown's pandas "NaN", and no dependency on markitdown's optional xlsx extra.
_XLSX_FIRST = {".xlsx", ".xlsm"}

_MAX_NESTED_EMAIL_DEPTH = 10

# The folder inside a project that the watcher converts. `add <project>` targets it.
INBOX_DIRNAME = "agent-inbox"
_LOG = logging.getLogger(__name__)


def _validate_config(data: object, base: Path) -> dict:
    if not isinstance(data, dict):
        raise ValueError("config must be a JSON object")
    cfg = dict(_DEFAULT_CONFIG)
    cfg.update(data)
    folders = cfg["watchFolders"]
    if not isinstance(folders, list) or any(
        not isinstance(p, str) or not p.strip() or "\0" in p for p in folders
    ):
        raise ValueError("watchFolders must be a list of non-empty folder paths")
    normalized = []
    for entry in folders:
        path = Path(entry).expanduser()
        normalized.append(str((path if path.is_absolute() else base / path).resolve()))
    cfg["watchFolders"] = list(dict.fromkeys(normalized))
    poll = cfg["pollSeconds"]
    try:
        valid_poll = not isinstance(poll, bool) and isinstance(poll, (int, float)) and \
            math.isfinite(poll) and poll >= 1
    except OverflowError:
        valid_poll = False
    if not valid_poll:
        raise ValueError("pollSeconds must be a finite number of at least 1")
    return cfg


def load_config(path: Path | None = None) -> dict:
    """Load and validate configuration; only a missing file uses defaults."""
    path = path if path is not None else CONFIG_PATH
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        data = {}
    except (ValueError, OSError) as exc:
        raise ValueError(f"Cannot read config {path}: {exc}") from exc
    try:
        return _validate_config(data, path.parent)
    except ValueError as exc:
        raise ValueError(f"Invalid config {path}: {exc}") from exc


def save_config(cfg: dict, path: Path | None = None) -> None:
    path = path if path is not None else CONFIG_PATH
    content = json.dumps(_validate_config(cfg, path.parent), indent=2) + "\n"
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent,
                                     prefix=".config-", delete=False) as fh:
        temp = Path(fh.name)
        try:
            fh.write(content)
            fh.flush()
            os.fsync(fh.fileno())
        except BaseException:
            temp.unlink()
            raise
    try:
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTS


# --------------------------------------------------------------------------- #
# HTML -> text
# --------------------------------------------------------------------------- #

def html_to_text(html: str) -> str:
    """Best-effort HTML to plain text. Uses html2text if installed."""
    if not html:
        return ""
    try:
        import html2text  # type: ignore

        h = html2text.HTML2Text()
        h.body_width = 0
        h.ignore_images = True
        return h.handle(html).strip()
    except ImportError:
        return _basic_html_to_text(html)


def _basic_html_to_text(html: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?</\1>", "", html)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</p\s*>", "\n\n", text)
    text = re.sub(r"(?i)</div\s*>", "\n", text)
    text = re.sub(r"(?i)</tr\s*>", "\n", text)
    text = re.sub(r"(?i)</h[1-6]\s*>", "\n\n", text)
    text = re.sub(r"(?is)<[^>]+>", "", text)
    text = _html.unescape(text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- #
# Filenames and Markdown rendering
# --------------------------------------------------------------------------- #

def slugify(text: str) -> str:
    slug = (text or "").strip().lower()
    slug = re.sub(r"^(re|fwd?|fw)\s*:\s*", "", slug)
    slug = re.sub(r"[^a-z0-9]+", "-", slug).strip("-")
    slug = slug[:60].strip("-")
    return slug or "untitled"


def _email_date_prefix(date_header: str | None, captured_iso: str) -> str:
    if date_header:
        try:
            dt = parsedate_to_datetime(date_header)
            if dt is not None:
                return dt.strftime("%Y-%m-%d")
        except (TypeError, ValueError, IndexError):
            pass
    return captured_iso[:10]


def build_filename(meta: dict) -> str:
    prefix = meta.get("date_prefix") or (meta.get("captured", "") or "")[:10] \
        or datetime.now().strftime("%Y-%m-%d")
    title = meta.get("subject") or meta.get("title") or "untitled"
    return f"{prefix}--{slugify(title)}.md"


def _yaml_str(value) -> str:
    return json.dumps("" if value is None else str(value), ensure_ascii=True)


def _human_size(num) -> str:
    try:
        num = float(num)
    except (TypeError, ValueError):
        return ""
    for unit in ("B", "KB", "MB", "GB"):
        if num < 1024 or unit == "GB":
            return f"{num:.0f} {unit}" if unit == "B" else f"{num:.1f} {unit}"
        num /= 1024
    return f"{num:.1f} GB"


def render_markdown(meta: dict, body: str) -> str:
    if meta.get("kind") == "email":
        return _render_email(meta, body)
    return _render_file(meta, body)


def _render_email(meta: dict, body_text: str) -> str:
    attachments = meta.get("attachments") or []
    att_names = [a.get("name", "attachment") for a in attachments]

    lines = ["---"]
    lines.append(f"subject: {_yaml_str(meta.get('subject'))}")
    lines.append(f"from: {_yaml_str(meta.get('from'))}")
    lines.append(f"to: {_yaml_str(meta.get('to'))}")
    if meta.get("cc"):
        lines.append(f"cc: {_yaml_str(meta.get('cc'))}")
    lines.append(f"date: {_yaml_str(meta.get('date'))}")
    if meta.get("message_id"):
        lines.append(f"message_id: {_yaml_str(meta.get('message_id'))}")
    if att_names:
        lines.append("attachments: [" + ", ".join(_yaml_str(n) for n in att_names) + "]")
    lines.append(f"captured: {_yaml_str(meta.get('captured'))}")
    lines.append(f"source: {_yaml_str(meta.get('source'))}")
    lines.append("---")
    lines.append("")
    lines.extend(_render_email_content(meta, body_text, heading_level=1))
    lines.append("")
    return "\n".join(lines)


def _render_email_content(meta: dict, body_text: str, heading_level: int) -> list[str]:
    lines = [f"{'#' * min(heading_level, 6)} {(meta.get('subject') or '(no subject)').strip()}", ""]
    lines.append(f"**From:** {meta.get('from') or ''}  ")
    lines.append(f"**To:** {meta.get('to') or ''}  ")
    if meta.get("cc"):
        lines.append(f"**Cc:** {meta.get('cc')}  ")
    lines.append(f"**Date:** {meta.get('date') or ''}")
    lines.extend(["", (body_text or "").strip()])

    attachments = meta.get("attachments") or []
    if not attachments:
        return lines

    lines.extend(["", "---", "", f"{'#' * min(heading_level + 1, 6)} Attachments"])
    for index, attachment in enumerate(attachments, start=1):
        lines.extend(["", f"{'#' * min(heading_level + 2, 6)} {index}. {attachment.get('name', 'attachment')}"])
        size = _human_size(attachment.get("size"))
        if size:
            lines.extend(["", f"**Size:** {size}"])

        error = attachment.get("error")
        attachment_meta = attachment.get("meta") or {}
        attachment_body = attachment.get("body") or ""
        if error:
            lines.extend(["", f"> **Conversion unavailable:** {error}"])
        elif attachment_meta.get("kind") == "email":
            lines.append("")
            lines.extend(_render_email_content(attachment_meta, attachment_body, heading_level + 3))
        else:
            lines.extend(["", attachment_body.strip()])

        if index < len(attachments):
            lines.extend(["", "---"])
    return lines


def _render_file(meta: dict, body: str) -> str:
    lines = ["---"]
    lines.append(f"title: {_yaml_str(meta.get('title'))}")
    lines.append(f"source_file: {_yaml_str(meta.get('source_file'))}")
    lines.append(f"type: {_yaml_str(meta.get('type'))}")
    lines.append(f"captured: {_yaml_str(meta.get('captured'))}")
    lines.append(f"source: {_yaml_str(meta.get('source'))}")
    lines.append("---")
    lines.append("")
    lines.append(f"# {(meta.get('title') or 'Untitled').strip()}")
    lines.append("")
    lines.append((body or "").strip())
    lines.append("")
    return "\n".join(lines)


def write_markdown(dest_dir: Path, meta: dict, body: str) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    return write_unique_text(dest_dir / build_filename(meta), render_markdown(meta, body))


def unique_paths(target: Path) -> Iterator[Path]:
    yield target
    counter = 2
    while True:
        yield target.with_name(f"{target.stem}-{counter}{target.suffix}")
        counter += 1


def write_unique_text(target: Path, content: str) -> Path:
    """Write output without replacing existing files or symlinks."""
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=target.parent,
                                     prefix=".automdconvert-", delete=False) as fh:
        temp = Path(fh.name)
        try:
            fh.write(content)
            fh.flush()
            os.fsync(fh.fileno())
        except BaseException:
            temp.unlink()
            raise
    try:
        return publish_unique_file(temp, target)
    finally:
        temp.unlink(missing_ok=True)


def publish_unique_file(source: Path, target: Path) -> Path:
    """Use atomic hard-link publication, or exclusive copying on other filesystems."""
    candidates = unique_paths(target)
    while True:
        candidate = next(candidates)
        try:
            os.link(source, candidate, follow_symlinks=False)
        except FileExistsError:
            continue
        except OSError as exc:
            if exc.errno not in {errno.ENOTSUP, errno.EPERM, errno.EXDEV, errno.ENOSYS}:
                raise
            try:
                output = candidate.open("xb")
            except FileExistsError:
                continue
            try:
                with output, source.open("rb") as input_file:
                    shutil.copyfileobj(input_file, output)
                    output.flush()
                    os.fsync(output.fileno())
                shutil.copystat(source, candidate)
            except BaseException:
                candidate.unlink()
                raise
        return candidate


# --------------------------------------------------------------------------- #
# Conversion dispatch
# --------------------------------------------------------------------------- #

def convert_file(path: Path, *, _email_depth: int = 0) -> tuple[dict, str]:
    """Turn a supported file into (meta, markdown_body). Raises on failure."""
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS | _TEXTUTIL_EXTS | _XLSX_FIRST:
        raise ValueError(f"unsupported file type: {ext or path.name}")
    captured = datetime.now(timezone.utc).isoformat(timespec="seconds")

    if ext == ".eml":
        meta, body = _eml_to_meta_body(path, _email_depth)
        meta["kind"] = "email"
        meta["date_prefix"] = _email_date_prefix(meta.get("date"), captured)
        return meta, body

    body = _convert_other(path)
    if not body or not body.strip():
        raise ValueError(f"no converter produced output for '{ext or path.name}'")

    try:
        date_prefix = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")
    except OSError:
        date_prefix = captured[:10]

    meta = {
        "kind": "file",
        "title": path.stem,
        "source_file": path.name,
        "type": ext.lstrip(".") or "file",
        "captured": captured,
        "date_prefix": date_prefix,
        "source": "drag-file",
    }
    return meta, body


def _convert_other(path: Path) -> str | None:
    ext = path.suffix.lower()

    if ext == ".txt":
        return _read_text(path)
    if ext in {".csv", ".json", ".xml"}:
        text = _read_text(path).rstrip()
        if not text:
            return None
        fence = "`" * max(3, 1 + max((len(m) for m in re.findall(r"`+", text)), default=0))
        return f"{fence}{ext.lstrip('.')}\n{text}\n{fence}"
    if ext in {".html", ".htm"}:
        return html_to_text(_read_text(path))

    # RTF/DOC/ODT: markitdown returns raw markup, so use textutil first.
    if ext in _TEXTUTIL_FIRST:
        t = _textutil(path)
        if t and t.strip():
            return t

    # XLSX/XLSM: openpyxl gives cleaner tables than markitdown (no "NaN" blanks).
    if ext in _XLSX_FIRST:
        x = _xlsx_to_markdown(path)
        if x and x.strip():
            return x

    md = _markitdown(path)
    if md and md.strip():
        return md

    if ext in _TEXTUTIL_EXTS:
        t = _textutil(path)
        if t and t.strip():
            return t

    return None


def _markitdown(path: Path) -> str | None:
    try:
        from markitdown import MarkItDown  # type: ignore
    except ImportError:
        return None
    try:
        return MarkItDown().convert(str(path)).text_content
    except Exception as exc:
        _LOG.warning("markitdown could not convert %s: %s", path.name, exc)
        return None


def _xlsx_to_markdown(path: Path) -> str | None:
    """Render an .xlsx/.xlsm workbook as one Markdown table per non-empty sheet."""
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        return None
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        _LOG.warning("openpyxl could not open %s: %s", path.name, exc)
        return None

    def cell(value) -> str:
        if value is None:
            return ""
        return str(value).replace("|", "\\|").replace("\n", " ").replace("\r", " ").strip()

    out: list[str] = []
    with closing(wb), closing(load_workbook(path, read_only=True, data_only=False)) as formulas:
        for ws in wb.worksheets:
            rows = []
            for values, expressions in zip(
                ws.iter_rows(values_only=True),
                formulas[ws.title].iter_rows(values_only=True),
            ):
                row = [
                    value if value is not None else
                    expression if isinstance(expression, str) and expression.startswith("=") else None
                    for value, expression in zip(values, expressions)
                ]
                if any(c is not None and str(c).strip() != "" for c in row):
                    rows.append(row)
            if not rows:
                continue
            width = max(len(r) for r in rows)
            out.append(f"## {ws.title}")
            header = list(rows[0]) + [None] * (width - len(rows[0]))
            out.append("| " + " | ".join(cell(c) for c in header) + " |")
            out.append("| " + " | ".join(["---"] * width) + " |")
            for r in rows[1:]:
                padded = list(r) + [None] * (width - len(r))
                out.append("| " + " | ".join(cell(c) for c in padded) + " |")
            out.append("")
    return "\n".join(out).strip() or None


def _textutil(path: Path) -> str | None:
    try:
        res = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            capture_output=True,
            text=True,
            timeout=120,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        _LOG.warning("textutil could not convert %s: %s", path.name, exc)
        return None
    if res.returncode:
        _LOG.warning("textutil could not convert %s: %s", path.name, res.stderr.strip())
    return res.stdout if res.returncode == 0 else None


def _read_text(path: Path) -> str:
    payload = path.read_bytes()
    if payload.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encoding = "utf-32"
    elif payload.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        encoding = "utf-8-sig"
        if path.suffix.lower() in {".html", ".htm", ".xml"} and not payload.startswith(b"\xef\xbb\xbf"):
            declaration = re.search(
                rb"""(?is)(?:<meta\b[^>]*\bcharset|<\?xml\b[^>]*\bencoding)\s*=\s*["']?\s*([\w.-]+)""",
                payload[:4096],
            )
            if declaration:
                encoding = declaration.group(1).decode("ascii")
    try:
        "".encode(encoding)
    except LookupError:
        _LOG.warning("%s declares unknown encoding %r; decoding as UTF-8", path.name, encoding)
        encoding = "utf-8-sig"
    try:
        return payload.decode(encoding)
    except UnicodeDecodeError:
        _LOG.warning("%s contains invalid %s; replacing undecodable characters", path.name, encoding)
        return payload.decode(encoding, errors="replace")


# --------------------------------------------------------------------------- #
# .eml parsing
# --------------------------------------------------------------------------- #

def _eml_to_meta_body(eml_path: Path, email_depth: int = 0) -> tuple[dict, str]:
    with open(eml_path, "rb") as fh:
        msg = BytesParser(policy=policy.default).parse(fh)

    def hdr(name: str) -> str:
        value = msg[name]
        return "" if value is None else str(value).strip()

    body_part = msg.get_body(preferencelist=("plain", "html"))

    body_text = ""
    if body_part is not None:
        ctype = body_part.get_content_type()
        content = _email_text(body_part)
        if ctype == "text/html":
            body_text = html_to_text(content)
        else:
            body_text = content

    attachments = []
    for part in _iter_email_attachments(msg):
        name = _attachment_name(part.get_filename())
        attachment = {"name": name, "size": None}
        try:
            payload = _attachment_bytes(part)
            if payload is None:
                raise ValueError("The attachment payload could not be read.")
            attachment["size"] = len(payload)
            suffix = Path(name).suffix.lower() or _attachment_suffix(part.get_content_type())
            if suffix == ".eml" and email_depth >= _MAX_NESTED_EMAIL_DEPTH:
                raise ValueError("Maximum nested email depth reached.")
            if part.get_content_maintype() == "text" and suffix in {
                ".txt", ".csv", ".json", ".xml", ".html", ".htm",
            }:
                payload = _email_text(part).encode("utf-8")
            with tempfile.TemporaryDirectory(prefix="automdconvert-") as temp_dir:
                attachment_path = Path(temp_dir) / f"attachment{suffix}"
                attachment_path.write_bytes(payload)
                attachment_meta, attachment_body = convert_file(
                    attachment_path,
                    _email_depth=email_depth + 1,
                )
            if attachment_meta.get("kind") == "file":
                attachment_meta.update({
                    "title": Path(name).stem or name,
                    "source_file": name,
                    "type": suffix.lstrip(".") or "file",
                    "source": "email-attachment",
                })
            else:
                attachment_meta["source"] = "email-attachment"
            attachment.update({"meta": attachment_meta, "body": attachment_body})
        except Exception as exc:
            attachment["error"] = str(exc) or "No converter produced Markdown."
        attachments.append(attachment)

    if not body_text.strip() and not attachments:
        raise ValueError("Email contains no readable body or attachments.")

    meta = {
        "subject": hdr("subject"),
        "from": hdr("from"),
        "to": hdr("to"),
        "cc": hdr("cc"),
        "date": hdr("date"),
        "message_id": hdr("message-id"),
        "attachments": attachments,
        "captured": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source": "drag-eml",
    }
    return meta, body_text


def _email_text(part) -> str:
    try:
        content = part.get_content()
    except LookupError:
        _LOG.warning("Unknown email charset %r; decoding as UTF-8", part.get_content_charset())
        return (part.get_payload(decode=True) or b"").decode("utf-8", errors="replace")
    if not isinstance(content, str):
        raise ValueError("Email text part did not contain text.")
    return content


def _iter_email_attachments(message):
    attachments = {id(part) for part in message.iter_attachments()}
    for part in message.iter_parts():
        if part.get_content_maintype() == "multipart" and \
                part.get_content_disposition() != "attachment" and not part.get_filename():
            yield from _iter_email_attachments(part)
        elif id(part) in attachments or part.get_content_disposition() == "attachment" or \
                part.get_filename():
            yield part


def _attachment_name(filename: str | None) -> str:
    name = (filename or "attachment").replace("\\", "/").split("/")[-1]
    return re.sub(r"[\r\n]+", " ", name).strip() or "attachment"


def _attachment_suffix(content_type: str) -> str:
    if content_type == "message/rfc822":
        return ".eml"
    return mimetypes.guess_extension(content_type, strict=False) or ""


def _attachment_bytes(part) -> bytes | None:
    payload = part.get_payload(decode=True)
    if isinstance(payload, bytes):
        return payload

    content = part.get_content()
    if isinstance(content, bytes):
        return content
    if isinstance(content, str):
        return content.encode(part.get_content_charset() or "utf-8", errors="replace")
    if hasattr(content, "as_bytes"):
        return content.as_bytes(policy=policy.default)

    raw_payload = part.get_payload()
    if isinstance(raw_payload, list):
        messages = [item.as_bytes(policy=policy.default) for item in raw_payload if hasattr(item, "as_bytes")]
        return b"\n".join(messages) if messages else None
    return None
